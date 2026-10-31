"""
Universal Query-Adaptive Extraction and Re-Ranking Service
==========================================================
Provides robust, domain-agnostic semantic re-ranking and adaptive sectioning:
1. Universal lexical & semantic chunk scoring (hybrid retrieval).
2. Never destructively discards valid semantic vector results.
3. Adaptive on-demand document sectioning for PDF, DOCX, CSV, Excel, TXT.
"""

import asyncio
import io
import re
import logging
from typing import List, Optional, Dict, Any

from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.future import select

from app.models import DocumentChunk
from app.services.extractor_service import chunk_text
from app.services.embedding_service import get_or_create_embeddings

logger = logging.getLogger(__name__)

SIMILARITY_THRESHOLD = 0.35

# Standard English stopwords
_STOPWORDS = {
    "what", "which", "where", "when", "why", "who", "how", "the", "and", "is",
    "in", "at", "of", "to", "a", "an", "for", "on", "with", "as", "by", "from",
    "it", "this", "that", "these", "those", "are", "was", "were", "be", "been",
    "have", "has", "had", "do", "does", "did", "can", "could", "should", "would",
    "may", "might", "must", "will", "shall", "me", "my", "you", "your", "them",
    "their", "he", "she", "his", "her", "i", "we", "us", "our", "about", "find",
    "search", "locate", "identify", "look", "show", "important",
    # Additional common words to prevent false keyword triggers on general queries
    "out", "any", "sort", "matter", "so", "up", "down", "over", "under", "all",
    "some", "every", "each", "other", "into", "through", "during", "before", "after",
    "above", "below", "between", "both", "few", "more", "most", "no", "nor", "not",
    "only", "own", "same", "than", "too", "very", "just", "now", "also", "well",
    "way", "even", "new", "want", "give", "tell", "say", "ask", "know", "please",
    "like", "get", "make", "take", "see", "come", "think", "use",
}

_QUERY_SYNONYMS = {
    "agreement": (
        "contract",
        "memorandum of understanding",
        "MOU",
        "NDA",
        "non-disclosure",
        "terms",
        "parties",
        "obligations",
    ),
    "contract": ("agreement", "MOU", "NDA", "terms", "parties", "obligations"),
}
_QUERY_TERM_ALIASES = {
    "aggrement": "agreement",
    "aggreement": "agreement",
}


def _extract_query_tokens(query: str) -> List[str]:
    """
    Parse query into meaningful keywords (length >= 2, skipping common stopwords).
    """
    query_lower = query.lower()
    tokens = re.findall(r"[a-zA-Z0-9_\-\.]+", query_lower)
    meaningful_tokens = [t for t in tokens if len(t) >= 2 and t not in _STOPWORDS]
    expanded_tokens = list(meaningful_tokens)

    for token in meaningful_tokens:
        normalized = token.rstrip("s")
        normalized = _QUERY_TERM_ALIASES.get(normalized, normalized)
        if normalized != token:
            expanded_tokens.append(normalized)
        expanded_tokens.extend(_QUERY_SYNONYMS.get(normalized, ()))

    return list(dict.fromkeys(expanded_tokens))


def expand_query_for_retrieval(query: str) -> str:
    """Add domain synonyms to improve recall without changing the user-facing query."""
    terms = _extract_query_tokens(query)
    return " ".join(dict.fromkeys([query, *terms]))

def score_chunk_relevance(chunk_text_str: str, query: str) -> int:
    """
    Returns an integer keyword-overlap score for a chunk vs the user query.
    Higher = more lexical overlap.
    """
    keywords = _extract_query_tokens(query)
    if not keywords or not chunk_text_str:
        return 0

    text_lower = chunk_text_str.lower()
    score = 0

    # Exact full query match bonus
    q_clean = query.strip().lower()
    if len(q_clean) > 4 and q_clean in text_lower:
        score += 8

    # Keyword hits
    for kw in keywords:
        if kw in text_lower:
            # Add points based on frequency (capped at 3 per keyword)
            count = min(text_lower.count(kw), 3)
            score += count

    return score

def chunk_relevance_filter(
    retrieved_chunks: List[Dict[str, Any]],
    query: str,
    min_score: int = 0,
) -> List[Dict[str, Any]]:
    """
    Re-ranks retrieved chunks using a hybrid formula:
    combined_score = (similarity * 0.70) + (normalized_keyword_score * 0.30)
    Never throws away high-similarity semantic chunks.
    """
    if not retrieved_chunks:
        return []

    scored = []
    max_kw = 1

    for ch in retrieved_chunks:
        kw_score = score_chunk_relevance(ch.get("chunk_text", ""), query)
        if kw_score > max_kw:
            max_kw = kw_score
        scored.append({**ch, "keyword_score": kw_score})

    # Hybrid ranking
    for ch in scored:
        sim = max(float(ch.get("similarity", 0.0)), 0.0)
        norm_kw = ch["keyword_score"] / max_kw
        ch["combined_score"] = round((sim * 0.70) + (norm_kw * 0.30), 4)

    # Sort descending by combined_score
    scored.sort(key=lambda c: c["combined_score"], reverse=True)
    return scored

def needs_adaptive_embedding(
    retrieved_chunks: List[Dict[str, Any]],
    query: str,
    sim_threshold: float = SIMILARITY_THRESHOLD,
) -> bool:
    """
    Returns True when vector search results are completely absent or very weak.
    """
    if not retrieved_chunks:
        return True

    # If filename or keyword match is strong, adaptive extraction is not needed
    if any(c.get("filename_matched") or c.get("keyword_score", 0) >= 2 for c in retrieved_chunks):
        return False

    top_similarity = max((c.get("similarity", 0.0) for c in retrieved_chunks), default=0.0)
    if top_similarity < 0.25:
        logger.info(f"[AdaptiveGate] Top similarity {top_similarity:.3f} < 0.25 — triggering adaptive extraction.")
        return True

    # Also trigger if top chunks have zero keyword overlap and low similarity
    top_kw = max((c.get("keyword_score", 0) for c in retrieved_chunks[:5]), default=0)
    if top_kw == 0 and top_similarity < 0.30:
        logger.info(f"[AdaptiveGate] No keyword hits and sim={top_similarity:.3f} < 0.30 — triggering adaptive.")
        return True

    return False

# ─────────────────────────────────────────────────────────────────────────────
# Universal Adaptive Section Extractors
# ─────────────────────────────────────────────────────────────────────────────

def _adaptive_sections_txt(raw_text: str, keywords: List[str]) -> List[str]:
    lines = raw_text.splitlines()
    if not lines:
        return []

    if not keywords:
        return [raw_text[:2000]]

    hit_indices = set()
    window = 25

    for idx, line in enumerate(lines):
        line_lower = line.lower()
        if any(kw in line_lower for kw in keywords):
            for w in range(max(0, idx - window), min(len(lines), idx + window + 1)):
                hit_indices.add(w)

    if not hit_indices:
        return [raw_text[:2000]]

    sorted_hits = sorted(hit_indices)
    blocks = []
    block_start = sorted_hits[0]
    prev = sorted_hits[0]
    for idx in sorted_hits[1:]:
        if idx - prev > 1:
            blocks.append("\n".join(lines[block_start: prev + 1]))
            block_start = idx
        prev = idx
    blocks.append("\n".join(lines[block_start: prev + 1]))
    return blocks

def _adaptive_sections_docx(docx_bytes: bytes, keywords: List[str]) -> List[str]:
    try:
        from docx import Document
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as e:
        logger.warning(f"DOCX adaptive read error: {e}")
        return []

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                paragraphs.append(row_text)

    if not paragraphs:
        return []

    if not keywords:
        return ["\n\n".join(paragraphs[:10])]

    hit_indices = set()
    window = 2
    for idx, para in enumerate(paragraphs):
        para_lower = para.lower()
        if any(kw in para_lower for kw in keywords):
            for w in range(max(0, idx - window), min(len(paragraphs), idx + window + 1)):
                hit_indices.add(w)

    if not hit_indices:
        return ["\n\n".join(paragraphs[:10])]

    sorted_hits = sorted(hit_indices)
    block_lines = [paragraphs[i] for i in sorted_hits]
    return ["\n\n".join(block_lines)]

def _adaptive_sections_csv(raw_bytes: bytes, keywords: List[str]) -> List[str]:
    try:
        import csv
        text = raw_bytes.decode("utf-8", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        rows = list(reader)
    except Exception:
        return _adaptive_sections_txt(raw_bytes.decode("utf-8", errors="ignore"), keywords)

    if not rows:
        return []

    header = rows[0]
    header_str = ",".join(header)
    matched_rows = [header_str]

    for row in rows[1:]:
        row_str = ",".join(str(cell) for cell in row)
        row_lower = row_str.lower()
        if not keywords or any(kw in row_lower for kw in keywords):
            matched_rows.append(row_str)

    # Return blocks
    return ["\n".join(matched_rows[:50])]

def _adaptive_sections_excel(excel_bytes: bytes, keywords: List[str]) -> List[str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes), data_only=True)
        blocks = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                continue
            header = [str(c) if c is not None else "" for c in rows[0]]
            header_str = f"[Sheet: {sheet_name}] " + " | ".join(header)
            matched_rows = [header_str]
            for row in rows[1:]:
                row_str = " | ".join(str(c) if c is not None else "" for c in row)
                row_lower = row_str.lower()
                if not keywords or any(kw in row_lower for kw in keywords):
                    matched_rows.append(row_str)
            if len(matched_rows) > 1:
                blocks.append("\n".join(matched_rows[:60]))
        return blocks
    except Exception as e:
        logger.warning(f"Excel adaptive read error: {e}")
        return []

def _adaptive_sections_pdf(pdf_bytes: bytes, keywords: List[str]) -> List[str]:
    """
    Extracts text from PDF pages matching query keywords.
    Falls back to first pages if no exact keywords found.
    """
    candidate_sections = []
    try:
        import fitz
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        matched_pages = []
        all_pages = []

        for page_idx, page in enumerate(doc):
            pt = (page.get_text() or "").strip()
            if pt:
                all_pages.append(f"[Page {page_idx + 1}]\n{pt}")
                if keywords and any(kw in pt.lower() for kw in keywords):
                    matched_pages.append(f"[Page {page_idx + 1}]\n{pt}")

        doc.close()
        if matched_pages:
            return matched_pages
        if all_pages:
            return all_pages[:8]
    except Exception:
        pass

    # Fallback with pypdf
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_bytes))
        all_pages = []
        matched_pages = []
        for page_idx, page in enumerate(reader.pages):
            text = (page.extract_text() or "").strip()
            if text:
                all_pages.append(f"[Page {page_idx + 1}]\n{text}")
                if keywords and any(kw in text.lower() for kw in keywords):
                    matched_pages.append(f"[Page {page_idx + 1}]\n{text}")
        if matched_pages:
            return matched_pages
        return all_pages[:8]
    except Exception as e:
        logger.warning(f"PDF adaptive section error: {e}")

    return []

# ─────────────────────────────────────────────────────────────────────────────
# Unified Adaptive Entry Point
# ─────────────────────────────────────────────────────────────────────────────

async def adaptive_embed_for_query(
    file_bytes: bytes,
    file_name: str,
    mime_type: str,
    query: str,
    db_file_id: Any,
    user_id: str,
    db: AsyncSession,
) -> List[str]:
    """
    Universal entry point: extracts and embeds only the sections of ANY document
    relevant to the query.
    """
    keywords = _extract_query_tokens(query)
    name_lower = file_name.lower()

    sections: List[str] = []

    if file_bytes[:4] == b"%PDF" or "pdf" in mime_type or name_lower.endswith(".pdf"):
        sections = await asyncio.to_thread(_adaptive_sections_pdf, file_bytes, keywords)
    elif name_lower.endswith(".docx") or "wordprocessingml" in mime_type:
        sections = await asyncio.to_thread(_adaptive_sections_docx, file_bytes, keywords)
    elif name_lower.endswith(".xlsx") or "spreadsheetml" in mime_type:
        sections = await asyncio.to_thread(_adaptive_sections_excel, file_bytes, keywords)
    elif name_lower.endswith(".csv") or "csv" in mime_type:
        sections = await asyncio.to_thread(_adaptive_sections_csv, file_bytes, keywords)
    else:
        raw_text = file_bytes.decode("utf-8", errors="ignore")
        sections = await asyncio.to_thread(_adaptive_sections_txt, raw_text, keywords)

    if not sections:
        return []

    # Chunk the sections
    all_chunks: List[str] = []
    for section in sections:
        all_chunks.extend(chunk_text(section, chunk_size=900, chunk_overlap=120))

    if not all_chunks:
        return []

    # Embed and save
    embeddings = await get_or_create_embeddings(all_chunks, db)

    max_idx_res = await db.execute(
        select(DocumentChunk.chunk_index)
        .where(DocumentChunk.file_id == db_file_id)
        .order_by(DocumentChunk.chunk_index.desc())
        .limit(1)
    )
    current_max = max_idx_res.scalar() or 0

    for rel_idx, (chunk_str, emb) in enumerate(zip(all_chunks, embeddings)):
        clean_chunk = chunk_str.replace("\x00", "")
        db_chunk = DocumentChunk(
            file_id=db_file_id,
            user_id=user_id,
            chunk_index=current_max + rel_idx + 1,
            chunk_text=clean_chunk,
            embedding=emb,
            token_count=len(clean_chunk.split()),
            metadata_={
                "query_targeted": True,
                "source": "adaptive_extraction",
                "query_hint": query,
            },
        )
        db.add(db_chunk)

    await db.commit()
    logger.info(f"[AdaptiveEmbed] Added {len(all_chunks)} targeted chunks for {file_name}.")
    return all_chunks

async def adaptive_embed_pdf_for_query(
    pdf_bytes: bytes,
    query: str,
    db_file_id: Any,
    user_id: str,
    db: AsyncSession,
) -> List[str]:
    """Backward-compatible alias."""
    return await adaptive_embed_for_query(
        file_bytes=pdf_bytes,
        file_name="document.pdf",
        mime_type="application/pdf",
        query=query,
        db_file_id=db_file_id,
        user_id=user_id,
        db=db,
    )
